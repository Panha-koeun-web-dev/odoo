# -*- coding: utf-8 -*-
import io
import openpyxl
from odoo.tests.common import TransactionCase
from odoo.addons.school_management.report.school_xlsx_report import SchoolXlsxReport


class TestSchoolXlsxReports(TransactionCase):

    def test_fee_report_xlsx_generation(self):
        """Test XLSX fee report generation and workbook structure."""
        data = SchoolXlsxReport.generate_fee_report(self.env)
        self.assertTrue(len(data) > 1000)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertIn('Fee Invoices', wb.sheetnames)
        self.assertIn('Financial Analytics', wb.sheetnames)

    def test_attendance_report_xlsx_generation(self):
        """Test XLSX attendance report generation and workbook structure."""
        data = SchoolXlsxReport.generate_attendance_report(self.env)
        self.assertTrue(len(data) > 1000)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertIn('Attendance Log', wb.sheetnames)
        self.assertIn('Student Summary', wb.sheetnames)

    def test_grade_report_xlsx_generation(self):
        """Test XLSX grade report generation and workbook structure."""
        data = SchoolXlsxReport.generate_grade_report(self.env)
        self.assertTrue(len(data) > 1000)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertIn('Exam Grades', wb.sheetnames)
        self.assertIn('Subject Analytics', wb.sheetnames)

    def test_student_report_xlsx_generation(self):
        """Test XLSX student roster report generation."""
        data = SchoolXlsxReport.generate_student_report(self.env)
        self.assertTrue(len(data) > 1000)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertIn('Students Roster', wb.sheetnames)

    def test_class_report_xlsx_generation(self):
        """Test XLSX class capacity report generation."""
        data = SchoolXlsxReport.generate_class_report(self.env)
        self.assertTrue(len(data) > 1000)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertIn('Classes Overview', wb.sheetnames)

    def test_export_xlsx_wizard_workflow(self):
        """Test export wizard creation, execution and state transition."""
        wizard = self.env['school.report.export.xlsx.wizard'].create({
            'report_type': 'fee',
        })
        self.assertEqual(wizard.state, 'choose')
        self.assertTrue(wizard.record_count >= 0)

        action = wizard.action_export_xlsx()
        self.assertEqual(wizard.state, 'get')
        self.assertTrue(wizard.file_data)
        self.assertTrue(wizard.file_name.endswith('.xlsx'))
        self.assertEqual(action.get('type'), 'ir.actions.act_url')
        self.assertIn('download=true', action.get('url'))
