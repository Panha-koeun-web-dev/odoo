import base64
import csv
import io
from datetime import datetime, date
import openpyxl

from odoo import fields, models, _
from odoo.exceptions import UserError, ValidationError


class ImportSchoolHolidayWizard(models.TransientModel):
    _name = 'school.import.holiday.wizard'
    _description = 'Import School Public Holidays'

    file = fields.Binary(
        string='File',
        required=True,
    )

    filename = fields.Char(
        string='Filename',
    )

    import_mode = fields.Selection(
        [
            ('create', 'Create New Holidays Only'),
            ('update', 'Create or Update Existing'),
        ],
        string='Import Mode',
        default='update',
        required=True,
    )

    auto_reschedule = fields.Boolean(
        string='Validate & Move Conflicting Schedules to Next Week',
        default=True,
        help="If checked, any timetable sessions scheduled on imported holiday dates will automatically be moved to the next week.",
    )

    def _parse_date_value(self, val):
        """Parse various date representations (date, datetime, string)."""
        if not val:
            return False
        if isinstance(val, date) and not isinstance(val, datetime):
            return val
        if isinstance(val, datetime):
            return val.date()

        val_str = str(val).strip()
        if not val_str:
            return False

        # Try common date formats
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%Y/%m/%d'):
            try:
                return datetime.strptime(val_str, fmt).date()
            except ValueError:
                pass

        try:
            parsed = fields.Date.from_string(val_str)
            if parsed:
                return parsed
        except Exception:
            pass

        raise ValidationError(
            _("Invalid date format: '%s'. Please use YYYY-MM-DD.") % val_str
        )

    def _read_rows(self, file_content, filename):
        """Extract list of row dicts from CSV or Excel file."""
        filename_lower = (filename or '').lower()
        rows = []

        if filename_lower.endswith(('.xlsx', '.xlsm', '.xltx')):
            try:
                wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
                ws = wb.active
                iter_rows = list(ws.iter_rows(values_only=True))
                if not iter_rows:
                    return rows
                headers = [str(h).strip().lower() if h is not None else '' for h in iter_rows[0]]
                for raw_row in iter_rows[1:]:
                    if not any(raw_row):
                        continue
                    row_dict = {}
                    for i, h in enumerate(headers):
                        if h and i < len(raw_row):
                            row_dict[h] = raw_row[i]
                    rows.append(row_dict)
            except Exception as e:
                raise UserError(_("Failed to read Excel file: %s") % str(e))
        else:
            # Assume CSV
            try:
                decoded = file_content.decode('utf-8-sig')
            except UnicodeDecodeError:
                try:
                    decoded = file_content.decode('latin1')
                except Exception as e:
                    raise UserError(_("Failed to decode CSV file: %s") % str(e))

            reader = csv.DictReader(io.StringIO(decoded))
            if reader.fieldnames:
                # normalize fieldnames to lowercase
                cleaned_fieldnames = {f: f.strip().lower() for f in reader.fieldnames if f}
                for r in reader:
                    cleaned_r = {cleaned_fieldnames[k]: v for k, v in r.items() if k in cleaned_fieldnames}
                    if any(cleaned_r.values()):
                        rows.append(cleaned_r)

        return rows

    def action_import(self):
        self.ensure_one()

        if not self.file:
            raise UserError(_('Please upload a CSV or Excel file.'))

        filename = self.filename or 'holidays.csv'
        filename_lower = filename.lower()

        if not (filename_lower.endswith('.csv') or filename_lower.endswith(('.xlsx', '.xlsm'))):
            raise UserError(
                _("Please upload a CSV (.csv) or Excel (.xlsx) file.\n\n"
                  "Columns required: name, date\n"
                  "Optional columns: end_date, holiday_type, description")
            )

        try:
            file_content = base64.b64decode(self.file)
        except Exception as error:
            raise UserError(_("Unable to decode uploaded file: %s") % error)

        rows = self._read_rows(file_content, filename)
        if not rows:
            raise UserError(_("The uploaded file contains no data rows."))

        Holiday = self.env['school.public.holiday'].with_context(skip_auto_reschedule=True)
        created = 0
        updated = 0
        errors = []
        affected_holidays = self.env['school.public.holiday']

        allowed_types = {'national', 'religious', 'school', 'other'}

        for row_number, row in enumerate(rows, start=2):
            try:
                name_val = row.get('name')
                name = str(name_val).strip() if name_val is not None else ''

                date_val = row.get('date')
                end_date_val = row.get('end_date')

                holiday_type_val = row.get('holiday_type')
                holiday_type = str(holiday_type_val).strip().lower() if holiday_type_val is not None else 'national'

                desc_val = row.get('description')
                description = str(desc_val).strip() if desc_val is not None else ''

                if not name:
                    raise ValidationError(_("Holiday name is required."))

                if not date_val:
                    raise ValidationError(_("Start date is required."))

                start_date = self._parse_date_value(date_val)
                end_date = self._parse_date_value(end_date_val) if end_date_val else False

                if holiday_type not in allowed_types:
                    holiday_type = 'national'

                values = {
                    'name': name,
                    'date': start_date,
                    'end_date': end_date,
                    'holiday_type': holiday_type,
                    'description': description,
                    'company_id': self.env.company.id,
                }

                existing = Holiday.search([
                    ('name', '=', name),
                    ('date', '=', start_date),
                    ('company_id', '=', self.env.company.id),
                ], limit=1)

                if existing and self.import_mode == 'update':
                    existing.write(values)
                    updated += 1
                    affected_holidays |= existing
                elif existing:
                    raise ValidationError(
                        _("Holiday '%s' on %s already exists.") % (name, start_date)
                    )
                else:
                    new_rec = Holiday.create(values)
                    created += 1
                    affected_holidays |= new_rec

            except Exception as error:
                errors.append(_("Row %d: %s") % (row_number, str(error)))

        if errors:
            message = (
                _("Import completed with errors:\n\n")
                + _("Created: %d\nUpdated: %d\nErrors: %d\n\n") % (created, updated, len(errors))
                + '\n'.join(errors[:20])
            )
            if len(errors) > 20:
                message += _("\n...and %d more errors.") % (len(errors) - 20)
            raise UserError(message)

        # Validate schedules and move to next week if enabled
        total_moved = 0
        if self.auto_reschedule and affected_holidays:
            for hol in affected_holidays:
                dates = hol.get_holiday_dates()
                moved = hol._move_timetable_sessions_for_dates(dates, holiday_name=hol.name)
                total_moved += moved

        notif_msg = _("%(created)d holiday(s) created, %(updated)d updated.") % {
            'created': created,
            'updated': updated,
        }
        if self.auto_reschedule:
            if total_moved > 0:
                notif_msg += _("\nValidated schedule: moved %(moved)d class session(s) to next week.") % {
                    'moved': total_moved,
                }
            else:
                notif_msg += _("\nValidated schedule: no conflicting sessions found.")

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Import Successful'),
                'message': notif_msg,
                'type': 'success',
                'sticky': False,
            },
        }
